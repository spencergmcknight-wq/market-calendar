"""
Forward P/E relay for the SPX Daily Market Analysis hub (runs once a day on GitHub).

  FactSet  — the weekly Earnings Insight report (PDF, published Thursdays/Fridays):
             "The forward 12-month P/E ratio for the S&P 500 is 22.4"
  S&P DJI  — the S&P 500 estimates spreadsheet, downloaded inside a real (headless)
             browser session: the next four quarters of operating EPS.

Results go to pe.json. If a source fails, its last good value is kept and the
reason is recorded, so the hub always has a number and knows how fresh it is.
"""
import json, re, io, os, datetime, html, urllib.request

OUT = "pe.json"
BROWSER = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36",
           "Accept": "text/html,application/pdf,application/xhtml+xml,*/*;q=0.8", "Accept-Language": "en-US,en;q=0.9"}
# tolerant: "forward 12-month P/E ratio for the S&P 500 is 22.4" with any hyphen / spacing
PE_RE = re.compile(r"forward\s*12\s*[\-\u2010-\u2015]?\s*month\s*P\s*/?\s*E\s*ratio\s*(?:for\s+the\s+S\s*&(?:amp;)?\s*P\s*500\s*(?:index\s*)?)?(?:is|was|of|stood\s+at|at)?\s*(\d{1,2}\.\d)", re.I)
FACTSET_PDF = "https://advantage.factset.com/hubfs/Website/Resources%20Section/Research%20Desk/Earnings%20Insight/EarningsInsight_{:%m%d%y}.pdf"
SPDJI_XLSX = "https://www.spglobal.com/spdji/en/documents/additional-material/sp-500-eps-est.xlsx"


def get(url, binary=False):
    req = urllib.request.Request(url, headers=BROWSER)
    with urllib.request.urlopen(req, timeout=30) as r:
        data = r.read()
        return data if binary else data.decode("utf-8", "replace")


def factset():
    from pypdf import PdfReader
    today, tried = datetime.date.today(), 0
    for back in range(0, 22):                       # the last three weeks' Thursdays and Fridays, newest first
        d = today - datetime.timedelta(days=back)
        if d.weekday() not in (3, 4):
            continue
        tried += 1
        try:
            raw = get(FACTSET_PDF.format(d), binary=True)
            text = " ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages[:4])
            m = PE_RE.search(re.sub(r"\s+", " ", text))
            if m:
                return float(m.group(1)), "Earnings Insight " + d.isoformat()
        except Exception:
            pass
    # backup: the Insight articles, matched on plain text
    page = get("https://insight.factset.com/topic/earnings")
    links = re.findall(r'href="((?:https://insight\.factset\.com)?/[a-z0-9\-]+)"', page, re.I)
    links = ["https://insight.factset.com" + l if l.startswith("/") else l for l in links]
    links = [l for i, l in enumerate(links) if l not in links[:i]]
    for url in links[:25]:
        try:
            text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", get(url))))
            m = PE_RE.search(text)
            if m:
                return float(m.group(1)), url
        except Exception:
            pass
    raise RuntimeError("not found (tried %d report dates and %d articles)" % (tried, min(25, len(links))))


def parse_spdji(raw):
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    col = None
    for r in rows:
        for j, v in enumerate(r):
            t = str(v or "")
            if re.search(r"OPERATING", t, re.I) and re.search(r"EARN|EPS", t, re.I) and re.search(r"SH", t, re.I):
                col = j
                break
        if col is not None:
            break
    if col is None:
        raise RuntimeError("operating EPS column not found")
    today, q = datetime.date.today(), {}
    for r in rows:
        d = r[0]
        if isinstance(d, datetime.datetime):
            d = d.date()
        elif isinstance(d, str):
            m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", d)
            d = datetime.date(int(m.group(3)), int(m.group(1)), int(m.group(2))) if m else None
        v = r[col] if col < len(r) else None
        if isinstance(d, datetime.date) and isinstance(v, (int, float)) and 0 < v < 200 and d > today and d not in q:
            q[d] = float(v)
    fut = sorted(q.items())[:4]
    if len(fut) < 4:
        raise RuntimeError("only %d future quarters found" % len(fut))
    return round(sum(v for _, v in fut), 2), [str(d) for d, _ in fut]


def spdji():
    try:                                            # plain download first (cheap)
        return parse_spdji(get(SPDJI_XLSX, binary=True))
    except Exception:
        pass
    from playwright.sync_api import sync_playwright   # then inside a real browser session
    with sync_playwright() as p:
        b = p.chromium.launch()
        ctx = b.new_context(user_agent=BROWSER["User-Agent"], accept_downloads=True)
        page = ctx.new_page()
        page.goto("https://www.spglobal.com/spdji/en/indices/equity/sp-500/", timeout=60000)
        page.wait_for_timeout(4000)
        resp = ctx.request.get(SPDJI_XLSX, timeout=60000)
        if resp.status != 200:
            b.close()
            raise RuntimeError("browser download HTTP %d" % resp.status)
        raw = resp.body()
        b.close()
    return parse_spdji(raw)


def main():
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    old = {}
    if os.path.exists(OUT):
        try:
            old = json.load(open(OUT))
        except Exception:
            old = {}
    out = {"updated_at": now, "ok": {}, "errors": {}, "factset": old.get("factset", {}), "spdji": old.get("spdji", {})}
    try:
        v, src = factset()
        out["factset"] = {"value": v, "source": src, "as_of": now}
        out["ok"]["factset"] = True
    except Exception as e:
        out["ok"]["factset"] = False
        out["errors"]["factset"] = str(e)[:200]
    try:
        eps, quarters = spdji()
        out["spdji"] = {"eps": eps, "quarters": quarters, "as_of": now}
        out["ok"]["spdji"] = True
    except Exception as e:
        out["ok"]["spdji"] = False
        out["errors"]["spdji"] = str(e)[:200]
    json.dump(out, open(OUT, "w"), indent=1)
    print(json.dumps(out, indent=1))


if __name__ == "__main__":
    main()
